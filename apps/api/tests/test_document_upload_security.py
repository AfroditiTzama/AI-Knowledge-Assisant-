from datetime import UTC, datetime
from io import BytesIO
from uuid import UUID, uuid4
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document as OpenXmlDocument
from fastapi import HTTPException, UploadFile
from pypdf import PdfWriter

from knowledge_assistant.application.documents.commands.process_document import (
    ProcessDocumentCommand,
)
from knowledge_assistant.application.documents.commands.upload_document import (
    UploadDocumentCommand,
)
from knowledge_assistant.application.documents.queries.list_documents import (
    ListDocumentsQuery,
)
from knowledge_assistant.domain.documents.entities import Document, DocumentStatus
from knowledge_assistant.domain.documents.file_storage import StoredFile
from knowledge_assistant.infrastructure.documents.text_extractors.docx_text_extractor import (
    DOCXTextExtractor,
)
from knowledge_assistant.infrastructure.documents.text_extractors.pdf_text_extractor import (
    PDFTextExtractor,
)
from knowledge_assistant.infrastructure.documents.text_extractors.txt_text_extractor import (
    TXTTextExtractor,
)
from knowledge_assistant.infrastructure.storage.local_file_storage import LocalFileStorage
from knowledge_assistant.presentation.api.v1.document_router import (
    _read_upload_bounded,
)


class InMemoryDocumentRepository:
    def __init__(self) -> None:
        self.documents: list[Document] = []
        self.updated_statuses: list[DocumentStatus] = []

    async def add(self, document: Document) -> Document:
        self.documents.append(document)
        return document

    async def update(self, document: Document) -> Document:
        self.updated_statuses.append(document.status)
        return document

    async def get_by_id(self, document_id: UUID) -> Document | None:
        return next(
            (document for document in self.documents if document.id == document_id),
            None,
        )

    async def list_by_owner_id(self, owner_id: UUID) -> list[Document]:
        return [document for document in self.documents if document.owner_id == owner_id]

    async def delete(self, document_id: UUID) -> None:
        self.documents = [
            document for document in self.documents if document.id != document_id
        ]


class RecordingFileStorage:
    def __init__(self, content: bytes = b"") -> None:
        self.content = content
        self.saved: list[str] = []

    async def save(self, *, file_content: bytes, original_filename: str) -> StoredFile:
        self.saved.append(original_filename)
        return StoredFile(
            stored_filename=f"stored-{original_filename}",
            storage_path=f"stored-{original_filename}",
            size_bytes=len(file_content),
        )

    async def read(self, storage_path: str) -> bytes:
        return self.content

    async def delete(self, storage_path: str) -> None:
        return None


def make_pdf(*, pages: int = 1) -> bytes:
    output = BytesIO()
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=100, height=100)
    writer.write(output)
    return output.getvalue()


def make_docx(text: str = "Knowledge") -> bytes:
    output = BytesIO()
    document = OpenXmlDocument()
    document.add_paragraph(text)
    document.save(output)
    return output.getvalue()


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("filename", "content_type", "content"),
    [
        ("knowledge.pdf", "application/pdf", make_pdf()),
        (
            "knowledge.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            make_docx(),
        ),
        ("knowledge.txt", "text/plain; charset=utf-8", b"Useful text"),
    ],
)
async def test_valid_supported_uploads_are_persisted(
    filename: str,
    content_type: str,
    content: bytes,
) -> None:
    repository = InMemoryDocumentRepository()
    storage = RecordingFileStorage()
    command = UploadDocumentCommand(repository, storage)

    document = await command.execute(
        owner_id=uuid4(),
        original_filename=filename,
        content_type=content_type,
        file_content=content,
    )

    assert storage.saved == [filename]
    assert document.status is DocumentStatus.UPLOADED


@pytest.mark.asyncio
async def test_unsupported_extension_is_rejected_before_storage() -> None:
    storage = RecordingFileStorage()
    command = UploadDocumentCommand(InMemoryDocumentRepository(), storage)

    with pytest.raises(ValueError, match="Unsupported document type"):
        await command.execute(
            owner_id=uuid4(),
            original_filename="payload.exe",
            content_type="application/octet-stream",
            file_content=b"payload",
        )

    assert storage.saved == []


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("filename", "content_type", "content", "message"),
    [
        ("report.pdf", "image/png", make_pdf(), "media type"),
        ("report.pdf", "application/pdf", b"not a pdf", "PDF signature"),
        (
            "report.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            b"not a docx",
            "DOCX signature",
        ),
        ("report.txt", "text/plain", make_pdf(), "TXT extension"),
    ],
)
async def test_mime_and_signature_mismatches_are_rejected_before_storage(
    filename: str,
    content_type: str,
    content: bytes,
    message: str,
) -> None:
    storage = RecordingFileStorage()
    command = UploadDocumentCommand(InMemoryDocumentRepository(), storage)

    with pytest.raises(ValueError, match=message):
        await command.execute(
            owner_id=uuid4(),
            original_filename=filename,
            content_type=content_type,
            file_content=content,
        )

    assert storage.saved == []


@pytest.mark.asyncio
async def test_upload_reader_stops_at_configured_limit() -> None:
    upload = UploadFile(file=BytesIO(b"123456"), filename="large.txt")

    with pytest.raises(HTTPException) as error:
        await _read_upload_bounded(upload, max_bytes=5, chunk_bytes=2)

    assert error.value.status_code == 413
    assert upload.file.tell() == 6


@pytest.mark.asyncio
async def test_malformed_pdf_and_docx_fail_safely() -> None:
    with pytest.raises(ValueError, match="invalid or corrupted"):
        await PDFTextExtractor().extract(
            file_content=b"%PDF-malformed",
            original_filename="broken.pdf",
            content_type="application/pdf",
        )

    with pytest.raises(ValueError, match="invalid or corrupted"):
        await DOCXTextExtractor().extract(
            file_content=b"PK\x03\x04malformed",
            original_filename="broken.docx",
            content_type=(
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ),
        )


@pytest.mark.asyncio
async def test_pdf_page_and_extracted_text_limits_are_enforced() -> None:
    with pytest.raises(ValueError, match="page limit"):
        await PDFTextExtractor(max_pages=1).extract(
            file_content=make_pdf(pages=2),
            original_filename="long.pdf",
            content_type="application/pdf",
        )

    with pytest.raises(ValueError, match="extracted text limit"):
        await TXTTextExtractor(max_text_characters=3).extract(
            file_content=b"four",
            original_filename="long.txt",
            content_type="text/plain",
        )


@pytest.mark.asyncio
async def test_docx_archive_expansion_limit_is_enforced_before_storage() -> None:
    output = BytesIO()
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "types")
        archive.writestr("word/document.xml", "x" * 2_000)

    storage = RecordingFileStorage()
    command = UploadDocumentCommand(
        InMemoryDocumentRepository(),
        storage,
        max_docx_uncompressed_bytes=1_000,
        max_docx_compression_ratio=10_000,
    )

    with pytest.raises(ValueError, match="decompressed size limit"):
        await command.execute(
            owner_id=uuid4(),
            original_filename="expanded.docx",
            content_type=(
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ),
            file_content=output.getvalue(),
        )

    assert storage.saved == []


@pytest.mark.asyncio
async def test_rejected_upload_leaves_no_stored_or_temporary_files(tmp_path) -> None:
    command = UploadDocumentCommand(
        InMemoryDocumentRepository(),
        LocalFileStorage(tmp_path),
    )

    with pytest.raises(ValueError, match="PDF signature"):
        await command.execute(
            owner_id=uuid4(),
            original_filename="invalid.pdf",
            content_type="application/pdf",
            file_content=b"invalid",
        )

    assert list(tmp_path.iterdir()) == []


@pytest.mark.asyncio
async def test_document_listing_remains_isolated_by_owner() -> None:
    first_owner = uuid4()
    second_owner = uuid4()
    repository = InMemoryDocumentRepository()
    command = UploadDocumentCommand(repository, RecordingFileStorage())

    first_document = await command.execute(
        owner_id=first_owner,
        original_filename="first.txt",
        content_type="text/plain",
        file_content=b"first",
    )
    await command.execute(
        owner_id=second_owner,
        original_filename="second.txt",
        content_type="text/plain",
        file_content=b"second",
    )

    documents = await ListDocumentsQuery(repository).execute(owner_id=first_owner)

    assert documents == [first_document]


@pytest.mark.asyncio
async def test_parser_failure_marks_document_failed() -> None:
    owner_id = uuid4()
    now = datetime.now(UTC)
    document = Document(
        id=uuid4(),
        owner_id=owner_id,
        original_filename="broken.pdf",
        stored_filename="broken.pdf",
        storage_path="broken.pdf",
        content_type="application/pdf",
        size_bytes=14,
        status=DocumentStatus.UPLOADED,
        created_at=now,
        updated_at=now,
    )
    repository = InMemoryDocumentRepository()
    repository.documents.append(document)
    unused_dependency = object()
    command = ProcessDocumentCommand(
        document_repository=repository,
        document_chunk_repository=unused_dependency,  # type: ignore[arg-type]
        file_storage=RecordingFileStorage(b"%PDF-malformed"),
        text_extractor=PDFTextExtractor(),
        text_chunker=unused_dependency,  # type: ignore[arg-type]
        embedding_service=unused_dependency,  # type: ignore[arg-type]
        vector_store=unused_dependency,  # type: ignore[arg-type]
    )

    with pytest.raises(ValueError, match="invalid or corrupted"):
        await command.execute(document_id=document.id, owner_id=owner_id)

    assert repository.updated_statuses == [
        DocumentStatus.PROCESSING,
        DocumentStatus.FAILED,
    ]
    assert document.status is DocumentStatus.FAILED
